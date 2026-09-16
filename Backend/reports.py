"""Relatórios produzidos a partir de uma leitura consistente do MySQL."""

import csv
from datetime import datetime, timezone
from io import BytesIO, StringIO
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Inches, Pt
from flask import Blueprint, g, send_file
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

from security import ApiError, database

reports = Blueprint("reports", __name__)


def money(value):
    return ("R$ " + f"{float(value or 0):,.2f}").replace(",", "_").replace(".", ",").replace("_", ".")


def report_data(kind):
    with database() as (_, cursor):
        if kind == "clientes":
            cursor.execute("SELECT razao_social, nome_fantasia, cnpj, estado, status, faturamento FROM clientes ORDER BY razao_social")
            rows = [[r["razao_social"], r["nome_fantasia"], r["cnpj"], r["estado"], r["status"], money(r["faturamento"])] for r in cursor.fetchall()]
            return "Clientes", ["Razão social", "Nome fantasia", "CNPJ", "UF", "Status", "Faturamento"], rows
        if kind == "produtos":
            cursor.execute("SELECT nome, categoria, marca, estoque, preco FROM produtos ORDER BY nome")
            rows = [[r["nome"], r["categoria"], r["marca"], r["estoque"], money(r["preco"])] for r in cursor.fetchall()]
            return "Produtos", ["Produto", "Categoria", "Marca", "Estoque", "Preço"], rows
        if kind == "financeiro":
            cursor.execute("SELECT COUNT(*) AS total, COALESCE(SUM(faturamento),0) AS revenue FROM clientes")
            clients = cursor.fetchone()
            cursor.execute("SELECT COUNT(*) AS total, COALESCE(SUM(estoque),0) AS stock, COALESCE(SUM(estoque*preco),0) AS value FROM produtos")
            products = cursor.fetchone()
            return "Indicadores financeiros", ["Indicador", "Valor"], [
                ["Clientes cadastrados", clients["total"]], ["Produtos cadastrados", products["total"]],
                ["Faturamento informado nos cadastros", money(clients["revenue"])],
                ["Unidades em estoque", products["stock"]], ["Valor do estoque a preço cadastrado", money(products["value"])],
            ]
    raise ApiError("Tipo de relatório inválido.", 404)


def make_pdf(title, columns, rows, author):
    stream = BytesIO()
    doc = SimpleDocTemplate(stream, pagesize=landscape(A4), rightMargin=30, leftMargin=30,
                            topMargin=30, bottomMargin=36, title=f"NexusGest — {title}", author=author)
    styles = getSampleStyleSheet()
    styles["Normal"].fontSize = 8
    styles["Normal"].leading = 11
    cell = lambda value: Paragraph(escape(str(value if value is not None else "")), styles["Normal"])
    heading = [cell(v) for v in columns]
    grid = [heading] + [[cell(v) for v in row] for row in rows]
    table = Table(grid, colWidths=[doc.width / len(columns)] * len(columns), repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dcefea")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f4f7f8")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
        ("LINEBELOW", (0, 0), (-1, 0), .5, colors.HexColor("#6ba596")),
    ]))
    generated = datetime.now(timezone.utc).strftime("%d/%m/%Y %H:%M UTC")
    story = [Paragraph("NexusGest", styles["Title"]), Paragraph(escape(title), styles["Heading2"]),
             cell(f"Gerado em {generated} por {author} · {len(rows)} registros"), Spacer(1, 16), table]
    if not rows:
        story.append(Paragraph("Nenhum registro cadastrado.", styles["Normal"]))
    def footer(canvas, document):
        canvas.setFont("Helvetica", 8)
        canvas.drawString(30, 18, "NexusGest | Dados da base no momento da emissão")
        canvas.drawRightString(landscape(A4)[0] - 30, 18, f"Página {document.page}")
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    stream.seek(0)
    return stream


def make_docx(title, columns, rows, author):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(11.7), Inches(8.3)
    section.left_margin = section.right_margin = Inches(.5)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_heading("NexusGest", 0)
    doc.add_heading(title, 1)
    doc.add_paragraph(f"Gerado por {author} · {datetime.now(timezone.utc):%d/%m/%Y %H:%M UTC}")
    table = doc.add_table(rows=1, cols=len(columns))
    table.style = "Light Shading Accent 1"
    for cell, value in zip(table.rows[0].cells, columns):
        cell.text = str(value)
    from docx.oxml import OxmlElement
    repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for cell, value in zip(table.add_row().cells, row):
            cell.text = str(value if value is not None else "")
    doc.add_paragraph(f"{len(rows)} registros. Valores informados nos cadastros; não constituem escrituração contábil.")
    stream = BytesIO()
    doc.save(stream)
    stream.seek(0)
    return stream


def make_csv(columns, rows):
    buffer = StringIO(newline="")
    writer = csv.writer(buffer, delimiter=";", lineterminator="\r\n")
    def safe(value):
        text = str(value if value is not None else "")
        return "'" + text if text.lstrip().startswith(("=", "+", "-", "@")) or text.startswith(("\t", "\r", "\n")) else text
    writer.writerow(columns)
    writer.writerows([[safe(v) for v in row] for row in rows])
    return BytesIO(buffer.getvalue().encode("utf-8-sig"))


@reports.get("/api/reports/<kind>.<extension>")
def download(kind, extension):
    if kind not in {"clientes", "produtos", "financeiro"} or extension not in {"pdf", "docx", "csv"}:
        raise ApiError("Relatório ou formato inválido.", 404)
    title, columns, rows = report_data(kind)
    if extension == "pdf":
        stream, mime = make_pdf(title, columns, rows, g.user["name"]), "application/pdf"
    elif extension == "docx":
        stream, mime = make_docx(title, columns, rows, g.user["name"]), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    else:
        stream, mime = make_csv(columns, rows), "text/csv; charset=utf-8"
    return send_file(stream, mimetype=mime, as_attachment=True, download_name=f"nexusgest-{kind}.{extension}")
